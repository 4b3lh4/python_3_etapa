from atividade_tabelas import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook()
F = "Arial"
hdr = PatternFill("solid", fgColor="1F4E78")

def cab(ws, cols, row=1):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row, j, c); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = hdr
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = 20
    ws.row_dimensions[row].height = 32

# --- Aba 1: comparação 2012 x 2026 (variações com fórmulas)
ws = wb.active; ws.title = "Comparação 2012-2026"
cols = ["Sigla","Código","Estado","Homens 2012 (%)","Mulheres 2012 (%)","Homens 2026 (%)","Mulheres 2026 (%)",
        "Var. Homens (p.p.)","Var. Mulheres (p.p.)","Maioria 2012","Maioria 2026"]
cab(ws, cols)
for i, r in enumerate(comp.itertuples(index=False), 2):
    vals = [r[0], r[1], r[2], r[3], r[4], r[5], r[6]]
    for j, v in enumerate(vals, 1): ws.cell(i, j, v)
    ws.cell(i, 8, f"=F{i}-D{i}"); ws.cell(i, 9, f"=G{i}-E{i}")
    ws.cell(i, 10, f'=IF(E{i}>D{i},"Mulheres","Homens")'); ws.cell(i, 11, f'=IF(G{i}>F{i},"Mulheres","Homens")')
n = len(comp) + 1
ws.cell(n+1, 3, "Média dos estados").font = Font(name=F, bold=True)
for col in "DEFGHI": ws[f"{col}{n+1}"] = f"=AVERAGE({col}2:{col}{n})"
ws.cell(n+2, 3, "Estados com maioria feminina").font = Font(name=F, bold=True)
ws[f"J{n+2}"] = f'=COUNTIF(J2:J{n},"Mulheres")'; ws[f"K{n+2}"] = f'=COUNTIF(K2:K{n},"Mulheres")'
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F, bold=c.font.bold)
        if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith("=")): 
            if c.column >= 4 and c.column <= 9: c.number_format = "0.0"
ws.freeze_panes = "D2"

# --- Aba 2: cruzamento com Tabela 1.1.1
w2 = wb.create_sheet("Cruzamento Tabela 1.1.1")
cols2 = list(final.columns)
cab(w2, cols2)
for i, r in enumerate(final.itertuples(index=False), 2):
    for j, v in enumerate(r, 1):
        c = w2.cell(i, j, v); c.font = Font(name=F)
        if isinstance(v, float): c.number_format = "0.00"
w2.freeze_panes = "D2"

# --- Aba 3: correlações
w3 = wb.create_sheet("Correlações")
w3.cell(1, 1, "Matriz de correlação (Pearson) - 27 UFs").font = Font(name=F, bold=True)
for j, c in enumerate(corr.columns, 2): 
    x = w3.cell(3, j, c); x.font = Font(name=F, bold=True); x.alignment = Alignment(wrap_text=True)
    w3.column_dimensions[get_column_letter(j)].width = 22
w3.column_dimensions["A"].width = 30
for i, (idx, row) in enumerate(corr.iterrows(), 4):
    w3.cell(i, 1, idx).font = Font(name=F, bold=True)
    for j, v in enumerate(row, 2):
        w3.cell(i, j, float(v)).font = Font(name=F)
w3.cell(11, 1, "Fonte: IBGE/PNAD Contínua (Tabela 1.1.1, 2022) e Tabelas 5 fornecidas na aula.").font = Font(name=F, italic=True)
wb.save("resultado_atividade.xlsx")
